"""
DocuMind — Autonomous Document-Generation Agent
===============================================
FastAPI + LangGraph + Groq agent that takes a natural-language business
request, autonomously plans the tasks, executes each step with an LLM,
self-critiques its draft, revises if needed, and produces a polished
.docx file. The React frontend is pre-built into frontend/dist and is
served as static SPA assets from the same FastAPI process, so this app
deploys as ONE Render service with no CORS or URL-swap headaches.
"""
from dotenv import load_dotenv
load_dotenv()
import os
import json
import uuid
import time
import logging
import re
from datetime import datetime
from typing import TypedDict, List, Optional

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, field_validator

from langgraph.graph import StateGraph, END
from groq import Groq
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# --------------------------------------------------------------------------
# Config / logging
# --------------------------------------------------------------------------

logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
logger = logging.getLogger("agent")

GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "")
GROQ_MODEL = os.environ.get("GROQ_MODEL", "llama3-70b-8192")

if not GROQ_API_KEY:
    logger.warning("GROQ_API_KEY not set. Set it as an environment variable before running.")

client = Groq(api_key=GROQ_API_KEY)

OUTPUT_DIR = "generated_docs"
os.makedirs(OUTPUT_DIR, exist_ok=True)

MAX_LLM_RETRIES = 5

# --------------------------------------------------------------------------
# LLM call wrapper with retry, sleep, and timeout
# --------------------------------------------------------------------------

def call_llm(system_prompt: str, user_prompt: str, json_mode: bool = False, temperature: float = 0.4) -> str:
    last_err = None
    for attempt in range(1, MAX_LLM_RETRIES + 1):
        try:
            kwargs = {}
            if json_mode:
                kwargs["response_format"] = {"type": "json_object"}
            resp = client.chat.completions.create(
                model=GROQ_MODEL,
                temperature=temperature,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                timeout=30,
                **kwargs,
            )
            return resp.choices[0].message.content
        except Exception as e:
            last_err = e
            logger.warning(f"LLM call attempt {attempt} failed: {e}")
            if attempt < MAX_LLM_RETRIES:
                time.sleep(2 * attempt)
    raise RuntimeError(f"LLM call failed after {MAX_LLM_RETRIES} attempts: {last_err}")


def safe_json_parse(raw: str, fallback: dict) -> dict:
    """Guardrail: LLMs sometimes wrap JSON in markdown fences or add stray text."""
    try:
        cleaned = raw.strip()
        if cleaned.startswith("```"):
            cleaned = cleaned.strip("`")
            if cleaned.lower().startswith("json"):
                cleaned = cleaned[4:]
        return json.loads(cleaned)
    except Exception as e:
        logger.error(f"JSON parse failed, using fallback. Error: {e}. Raw: {raw[:300]}")
        return fallback


# --------------------------------------------------------------------------
# Agent State
# --------------------------------------------------------------------------

class AgentState(TypedDict):
    request: str
    doc_type: str
    assumptions: List[str]
    task_list: List[dict]
    sections: List[dict]
    draft_sections: List[dict]
    critique: dict
    revised: bool
    file_path: str
    error: Optional[str]


# --------------------------------------------------------------------------
# Graph Nodes
# --------------------------------------------------------------------------

def node_classify_and_plan(state: AgentState) -> AgentState:
    system = (
        "You are an autonomous planning agent for a business document generator. "
        "Given a user's natural language request, decide: "
        "1) the most appropriate business document type "
        "(proposal, meeting_minutes, project_plan, business_report, technical_design, sop, product_spec), "
        "2) a logical ordered list of sections needed for that document type, "
        "3) any reasonable assumptions you must make because the request is ambiguous, "
        "incomplete, or has conflicting requirements. Always make sensible assumptions "
        "rather than refusing - state them explicitly. "
        "Respond ONLY with strict JSON: "
        '{"doc_type": "...", "title": "...", "assumptions": ["..."], '
        '"sections": ["Section Name 1", "Section Name 2", ...]}'
    )
    raw = call_llm(system, f"User request: {state['request']}", json_mode=True, temperature=0.3)
    fallback = {
        "doc_type": "business_report",
        "title": "Generated Business Document",
        "assumptions": ["Request was generic; defaulted to a standard business report structure."],
        "sections": ["Executive Summary", "Background", "Key Points", "Recommendations", "Next Steps"],
    }
    plan = safe_json_parse(raw, fallback)

    task_list = [{"step": i + 1, "title": f"Draft section: {s}", "status": "pending"}
                 for i, s in enumerate(plan.get("sections", fallback["sections"]))]
    task_list.insert(0, {"step": 0, "title": f"Classify request -> {plan.get('doc_type', 'business_report')}", "status": "done"})
    task_list.append({"step": len(task_list), "title": "Self-critique draft", "status": "pending"})
    task_list.append({"step": len(task_list) + 1, "title": "Export to DOCX", "status": "pending"})

    state["doc_type"] = plan.get("doc_type", "business_report")
    state["assumptions"] = plan.get("assumptions", [])
    state["task_list"] = task_list
    state["sections"] = [{"heading": h, "content": ""} for h in plan.get("sections", fallback["sections"])]
    state["_title"] = plan.get("title", "Generated Business Document")  # type: ignore
    logger.info(f"Planned doc_type={state['doc_type']} sections={[s['heading'] for s in state['sections']]}")
    return state


def node_execute_steps(state: AgentState) -> AgentState:
    assumptions_text = "\n".join(f"- {a}" for a in state["assumptions"]) or "None"
    drafted = []
    for sec in state["sections"]:
        system = (
            f"You are drafting the '{sec['heading']}' section of a {state['doc_type'].replace('_', ' ')} "
            "for a professional business audience. Write clear, concrete, well-structured content "
            "(use realistic mock data/numbers/names where specifics are missing - never leave placeholders "
            "like [insert here]). 120-220 words. Plain text, no markdown symbols."
        )
        user = (
            f"Original user request: {state['request']}\n\n"
            f"Assumptions already made by the planning stage:\n{assumptions_text}\n\n"
            f"Write the content for the section titled: {sec['heading']}"
        )
        content = call_llm(system, user, temperature=0.5)
        drafted.append({"heading": sec["heading"], "content": content.strip()})

    for t in state["task_list"]:
        if t["title"].startswith("Draft section"):
            t["status"] = "done"

    state["draft_sections"] = drafted
    logger.info(f"Executed {len(drafted)} section drafts.")
    return state


def node_reflect(state: AgentState) -> AgentState:
    joined = "\n\n".join(f"### {s['heading']}\n{s['content']}" for s in state["draft_sections"])
    system = (
        "You are a critical editor reviewing an AI-generated business document draft. "
        "Check it against the ORIGINAL user request. Identify concrete issues: missing "
        "requirements, unaddressed ambiguity, weak/generic content, inconsistent assumptions. "
        "Respond ONLY with strict JSON: "
        '{"needs_revision": true/false, "issues": ["issue 1", "issue 2"]}. '
        "If the draft is genuinely solid, set needs_revision to false with an empty issues list."
    )
    user = f"Original request: {state['request']}\n\nDraft:\n{joined}"
    raw = call_llm(system, user, json_mode=True, temperature=0.2)
    critique = safe_json_parse(raw, {"needs_revision": False, "issues": []})

    for t in state["task_list"]:
        if t["title"] == "Self-critique draft":
            t["status"] = "done"

    state["critique"] = critique
    logger.info(f"Reflection result: {critique}")
    return state


def node_revise(state: AgentState) -> AgentState:
    issues = state["critique"].get("issues", [])
    joined = "\n\n".join(f"### {s['heading']}\n{s['content']}" for s in state["draft_sections"])
    system = (
        "You are revising a business document draft to fix specific issues raised by an editor. "
        "Keep the same section headings and overall structure. Improve content quality, "
        "address every issue listed, and keep it professional and concrete. "
        'Respond ONLY with strict JSON: {"sections": [{"heading": "...", "content": "..."}, ...]}'
    )
    user = (
        f"Original request: {state['request']}\n\n"
        f"Issues to fix:\n" + "\n".join(f"- {i}" for i in issues) +
        f"\n\nCurrent draft:\n{joined}"
    )
    raw = call_llm(system, user, json_mode=True, temperature=0.4)
    fallback = {"sections": state["draft_sections"]}
    revised = safe_json_parse(raw, fallback)
    state["draft_sections"] = revised.get("sections", state["draft_sections"])
    state["revised"] = True
    logger.info("Draft revised based on self-critique.")
    return state


def node_export_docx(state: AgentState) -> AgentState:
    for t in state["task_list"]:
        if t["title"] == "Export to DOCX":
            t["status"] = "done"

    doc = Document()

    title = doc.add_heading(getattr_or(state, "_title", "Generated Business Document"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Document Type: {state['doc_type'].replace('_', ' ').title()}  |  Generated: {datetime.now().strftime('%B %d, %Y')}")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph()

    if state.get("assumptions"):
        doc.add_heading("Assumptions Made by the Agent", level=1)
        for a in state["assumptions"]:
            doc.add_paragraph(a, style="List Bullet")
        doc.add_paragraph()

    for sec in state["draft_sections"]:
        doc.add_heading(sec["heading"], level=1)
        for para in sec["content"].split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip())

    doc.add_page_break()
    doc.add_heading("Agent Execution Log", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Step", "Task", "Status"
    for t in state["task_list"]:
        row = table.add_row().cells
        row[0].text = str(t["step"])
        row[1].text = t["title"]
        row[2].text = t["status"]

    fname = f"{state['doc_type']}_{uuid.uuid4().hex[:8]}.docx"
    fpath = os.path.join(OUTPUT_DIR, fname)
    doc.save(fpath)

    state["file_path"] = fpath
    logger.info(f"Document exported to {fpath}")
    return state


def getattr_or(state: dict, key: str, default: str) -> str:
    return state.get(key, default)  # type: ignore


def route_after_reflect(state: AgentState) -> str:
    if state["critique"].get("needs_revision") and state["critique"].get("issues"):
        return "revise"
    return "export"


# --------------------------------------------------------------------------
# Build LangGraph
# --------------------------------------------------------------------------

def build_graph():
    graph = StateGraph(AgentState)
    graph.add_node("plan", node_classify_and_plan)
    graph.add_node("execute", node_execute_steps)
    graph.add_node("reflect", node_reflect)
    graph.add_node("revise", node_revise)
    graph.add_node("export", node_export_docx)

    graph.set_entry_point("plan")
    graph.add_edge("plan", "execute")
    graph.add_edge("execute", "reflect")
    graph.add_conditional_edges("reflect", route_after_reflect, {"revise": "revise", "export": "export"})
    graph.add_edge("revise", "export")
    graph.add_edge("export", END)

    return graph.compile()


agent_graph = build_graph()


# --------------------------------------------------------------------------
# Helpers to build response
# --------------------------------------------------------------------------

def _build_initial_state(request_text: str) -> AgentState:
    return {
        "request": request_text,
        "doc_type": "",
        "assumptions": [],
        "task_list": [],
        "sections": [],
        "draft_sections": [],
        "critique": {},
        "revised": False,
        "file_path": "",
        "error": None,
    }


def _build_response_from_state(state: AgentState) -> dict:
    fname = os.path.basename(state["file_path"]) if state.get("file_path") else ""
    return {
        "message": f"Successfully generated a {state['doc_type'].replace('_', ' ')} document.",
        "doc_type": state["doc_type"],
        "assumptions": state["assumptions"],
        "task_list": state["task_list"],
        "revised": state["revised"],
        "download_url": f"/download/{fname}" if fname else "",
    }


# --------------------------------------------------------------------------
# FastAPI app + static SPA serving
# --------------------------------------------------------------------------

app = FastAPI(title="DocuMind", version="3.0.0")

ALLOWED_ORIGINS = os.environ.get(
    "ALLOWED_ORIGINS",
    "http://localhost:5173,http://127.0.0.1:5173,http://localhost:3000"
).split(",")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[o.strip() for o in ALLOWED_ORIGINS if o.strip()] + ["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class AgentRequest(BaseModel):
    request: str = Field(..., min_length=3, max_length=4000)

    @field_validator("request")
    @classmethod
    def not_blank(cls, v: str) -> str:
        if not v.strip():
            raise ValueError("request must not be empty or whitespace only")
        return v.strip()


class AgentResponse(BaseModel):
    message: str
    doc_type: str
    assumptions: List[str]
    task_list: List[dict]
    revised: bool
    download_url: str


@app.post("/agent", response_model=AgentResponse)
def run_agent(payload: AgentRequest):
    initial_state = _build_initial_state(payload.request)
    try:
        final_state = agent_graph.invoke(initial_state)
    except Exception as e:
        logger.exception("Agent execution failed")
        raise HTTPException(status_code=500, detail=f"Agent execution failed: {str(e)}")
    return AgentResponse(**_build_response_from_state(final_state))


@app.get("/")
def root():
    dist_index = os.path.join(os.path.dirname(os.path.abspath(__file__)), "frontend", "dist", "index.html")
    if os.path.exists(dist_index):
        return FileResponse(dist_index)
    return RedirectResponse(url="/docs")


@app.get("/download/{filename}")
def download_file(filename: str):
    fpath = os.path.join(OUTPUT_DIR, filename)
    if not os.path.exists(fpath):
        raise HTTPException(status_code=404, detail="File not found")
        
    # Strip the hex suffix, replace underscores with spaces, and apply Title Case
    clean_name = re.sub(r'_[a-f0-9]+\.docx$', '', filename)
    clean_name = clean_name.replace('_', ' ').title() + '.docx'

    return FileResponse(
        fpath,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=clean_name,
    )


@app.get("/health")
def health():
    return {"status": "ok", "model": GROQ_MODEL}


_FRONTEND_DIST = os.path.join(os.path.dirname(os.path.abspath(__file__)), "frontend", "dist")
if os.path.isdir(_FRONTEND_DIST):
    app.mount("/assets", StaticFiles(directory=os.path.join(_FRONTEND_DIST, "assets")), name="assets")
    logger.info(f"Frontend assets mounted from {_FRONTEND_DIST}")
else:
    logger.warning(
        f"No frontend build found at {_FRONTEND_DIST}. Run `cd frontend && npm install && npm run build` "
        "to generate it, or access /docs for the API."
    )
